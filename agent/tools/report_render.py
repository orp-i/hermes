from __future__ import annotations

import re
import time
import os
import base64
from pathlib import Path
from typing import Any

import httpx
from xml.sax.saxutils import escape

from ..sandbox import SANDBOX_WORKSPACE_ROOT

SCHEMA = {
    "type": "function",
    "function": {
        "name": "render_report",
        "description": (
            "Render a Markdown report in the current session workspace into downloadable PDF and DOCX files. "
            "Use this after creating a final client-ready report .md in outputs/ so users receive both rendered documents. "
            "The tool also returns CRO consulting deliverable quality warnings that should be addressed before final delivery."
        ),
        "parameters": {
            "type": "object",
            "properties": {
                "markdown_path": {
                    "type": "string",
                    "description": "Markdown path relative to the session root, e.g. outputs/report.md.",
                },
                "output_basename": {
                    "type": "string",
                    "description": "Optional basename for generated files, without extension.",
                },
                "title": {
                    "type": "string",
                    "description": "Optional report title used in document metadata.",
                },
            },
            "required": ["markdown_path"],
        },
    },
}

REVIEW_SCHEMA = {
    "type": "function",
    "function": {
        "name": "review_report_layout",
        "description": (
            "Render PDF report pages to images and ask the configured multimodal model for layout feedback. "
            "Use this for CRO consulting report layout and content-readiness optimization iterations after render_report."
        ),
        "parameters": {
            "type": "object",
            "properties": {
                "pdf_path": {
                    "type": "string",
                    "description": "PDF path relative to the session root, e.g. outputs/report.pdf.",
                },
                "focus": {
                    "type": "string",
                    "description": "Optional layout review focus, e.g. readability, tables, hierarchy, spacing.",
                },
                "max_pages": {
                    "type": "integer",
                    "description": "Number of pages to review from the start of the PDF, default 2, max 4.",
                    "minimum": 1,
                    "maximum": 4,
                },
            },
            "required": ["pdf_path"],
        },
    },
}


def _session_root(session_id: str) -> Path:
    return Path(SANDBOX_WORKSPACE_ROOT) / "sessions" / session_id


def _resolve_session_path(session_id: str, rel_path: str) -> Path:
    root = _session_root(session_id).resolve()
    path = (root / rel_path).resolve()
    if path != root and root not in path.parents:
        raise ValueError("path must stay inside the session workspace")
    return path


def _safe_basename(value: str) -> str:
    value = Path(value).stem
    value = re.sub(r"[^A-Za-z0-9._\-\u4e00-\u9fff]+", "_", value).strip("._-")
    return value or "report"


def _parse_blocks(markdown: str) -> list[tuple[str, Any]]:
    blocks: list[tuple[str, Any]] = []
    code_mode = False
    code_lines: list[str] = []
    table_lines: list[str] = []

    def flush_table() -> None:
        nonlocal table_lines
        if table_lines:
            blocks.append(("table", table_lines))
            table_lines = []

    for raw in markdown.replace("\r\n", "\n").splitlines():
        line = raw.rstrip()
        if line.strip().startswith("```"):
            flush_table()
            if code_mode:
                blocks.append(("code", "\n".join(code_lines)))
                code_lines = []
                code_mode = False
            else:
                code_mode = True
            continue
        if code_mode:
            code_lines.append(raw)
            continue
        if re.fullmatch(r"\s*[-*_]{3,}\s*", line):
            flush_table()
            blocks.append(("rule", None))
            continue
        if "|" in line and line.strip().startswith("|") and line.strip().endswith("|"):
            table_lines.append(line)
            continue
        flush_table()
        if not line.strip():
            blocks.append(("space", None))
            continue
        heading = re.match(r"^(#{1,4})\s+(.+)$", line)
        if heading:
            blocks.append((f"h{len(heading.group(1))}", heading.group(2).strip()))
            continue
        bullet = re.match(r"^\s*[-*]\s+(.+)$", line)
        if bullet:
            blocks.append(("bullet", bullet.group(1).strip()))
            continue
        ordered = re.match(r"^\s*\d+[.)]\s+(.+)$", line)
        if ordered:
            blocks.append(("number", ordered.group(1).strip()))
            continue
        blocks.append(("para", line.strip()))
    flush_table()
    if code_mode and code_lines:
        blocks.append(("code", "\n".join(code_lines)))
    return blocks


def _strip_inline(text: str) -> str:
    text = re.sub(r"`([^`]+)`", r"\1", text)
    text = re.sub(r"\*\*([^*]+)\*\*", r"\1", text)
    text = re.sub(r"\*([^*]+)\*", r"\1", text)
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1 (\2)", text)
    return text


def _wrap_long_sequences(text: str, width: int = 80) -> str:
    wrapped_lines: list[str] = []
    for line in str(text).splitlines():
        if re.fullmatch(r"[A-Z*]{90,}", line.strip()):
            seq = line.strip()
            wrapped_lines.extend(seq[i : i + width] for i in range(0, len(seq), width))
        else:
            wrapped_lines.append(line)
    return "\n".join(wrapped_lines)


def _pdf_text(text: str) -> str:
    return escape(_strip_inline(text)).replace("\n", "<br/>")


def _table_rows(lines: list[str]) -> list[list[str]]:
    rows = []
    for line in lines:
        cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
        if cells and all(re.fullmatch(r":?-{3,}:?", c or "") for c in cells):
            continue
        rows.append([_strip_inline(c) for c in cells])
    return rows


def _render_docx(blocks: list[tuple[str, Any]], path: Path, title: str) -> None:
    from docx import Document
    from docx.enum.text import WD_BREAK
    from docx.shared import Pt
    from docx.oxml.ns import qn

    doc = Document()
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Pt(54)
    section.left_margin = section.right_margin = Pt(54)
    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    styles["Normal"].font.size = Pt(10.5)
    if title:
        doc.core_properties.title = title

    for kind, value in blocks:
        if kind.startswith("h"):
            level = min(int(kind[1]), 4)
            doc.add_heading(_strip_inline(str(value)), level=level)
        elif kind == "bullet":
            doc.add_paragraph(_strip_inline(str(value)), style="List Bullet")
        elif kind == "number":
            doc.add_paragraph(_strip_inline(str(value)), style="List Number")
        elif kind == "code":
            p = doc.add_paragraph()
            run = p.add_run(_wrap_long_sequences(str(value), 80))
            run.font.name = "Courier New"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Courier New")
            run.font.size = Pt(8.5)
        elif kind == "table":
            rows = _table_rows(value)
            if not rows:
                continue
            table = doc.add_table(rows=len(rows), cols=max(len(r) for r in rows))
            table.style = "Table Grid"
            table.autofit = True
            for i, row in enumerate(rows):
                for j, cell in enumerate(row):
                    table.cell(i, j).text = cell
        elif kind == "space":
            doc.add_paragraph("")
        elif kind == "rule":
            doc.add_paragraph("")
        else:
            doc.add_paragraph(_strip_inline(str(value)))
    footer = doc.sections[0].footer.paragraphs[0]
    footer.text = "Generated by Hermes"
    footer.runs[0].add_break(WD_BREAK.LINE)
    doc.save(path)


def _render_pdf(blocks: list[tuple[str, Any]], path: Path, title: str) -> None:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.pdfbase import pdfmetrics
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, LongTable, TableStyle, Preformatted, PageBreak

    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    styles = getSampleStyleSheet()
    base = ParagraphStyle(
        "HermesBase",
        parent=styles["BodyText"],
        fontName="STSong-Light",
        fontSize=10,
        leading=15,
        alignment=TA_LEFT,
        spaceAfter=5,
        wordWrap="CJK",
    )
    heading = {
        "h1": ParagraphStyle("H1", parent=base, fontSize=18, leading=24, spaceBefore=10, spaceAfter=8, keepWithNext=True),
        "h2": ParagraphStyle("H2", parent=base, fontSize=15, leading=20, spaceBefore=8, spaceAfter=6, keepWithNext=True),
        "h3": ParagraphStyle("H3", parent=base, fontSize=12, leading=17, spaceBefore=6, spaceAfter=4, keepWithNext=True),
        "h4": ParagraphStyle("H4", parent=base, fontSize=11, leading=16, spaceBefore=5, spaceAfter=3),
    }
    code_style = ParagraphStyle("Code", parent=base, fontName="Courier", fontSize=8, leading=10)
    cell_style = ParagraphStyle(
        "Cell",
        parent=base,
        fontSize=7.5,
        leading=10,
        spaceAfter=0,
        wordWrap="CJK",
    )
    doc = SimpleDocTemplate(
        str(path),
        pagesize=A4,
        rightMargin=18 * mm,
        leftMargin=18 * mm,
        topMargin=18 * mm,
        bottomMargin=18 * mm,
        title=title or None,
    )
    story = []
    for kind, value in blocks:
        if kind in heading:
            story.append(Paragraph(_pdf_text(str(value)), heading[kind]))
        elif kind == "bullet":
            story.append(Paragraph("• " + _pdf_text(str(value)), base))
        elif kind == "number":
            story.append(Paragraph(_pdf_text(str(value)), base))
        elif kind == "code":
            story.append(Preformatted(_wrap_long_sequences(str(value), 80), code_style))
        elif kind == "table":
            rows = _table_rows(value)
            if rows:
                max_cols = max(len(r) for r in rows)
                normalized = [row + [""] * (max_cols - len(row)) for row in rows]
                table_data = [[Paragraph(_pdf_text(str(cell)), cell_style) for cell in row] for row in normalized]
                available_width = A4[0] - (36 * mm)
                col_widths = [available_width / max_cols] * max_cols
                table = LongTable(table_data, repeatRows=1, colWidths=col_widths, splitByRow=True)
                table.setStyle(
                    TableStyle(
                        [
                            ("FONTNAME", (0, 0), (-1, -1), "STSong-Light"),
                            ("FONTSIZE", (0, 0), (-1, -1), 7.5),
                            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eef2f7")),
                            ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#c9d1d9")),
                            ("VALIGN", (0, 0), (-1, -1), "TOP"),
                            ("LEFTPADDING", (0, 0), (-1, -1), 4),
                            ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                        ]
                    )
                )
                story.append(table)
                story.append(Spacer(1, 5))
        elif kind == "space":
            story.append(Spacer(1, 4))
        elif kind == "rule":
            story.append(Spacer(1, 6))
        else:
            story.append(Paragraph(_pdf_text(str(value)), base))

    def _footer(canvas, doc_obj):
        canvas.saveState()
        canvas.setFont("STSong-Light", 8)
        canvas.setFillColor(colors.HexColor("#6b7280"))
        canvas.drawString(18 * mm, 10 * mm, "Generated by Hermes")
        canvas.drawRightString(A4[0] - 18 * mm, 10 * mm, f"Page {doc_obj.page}")
        canvas.restoreState()

    doc.build(story, onFirstPage=_footer, onLaterPages=_footer)


def _report_quality_warnings(markdown: str) -> list[str]:
    text = markdown.lower()
    warnings: list[str] = []
    required = [
        ("executive summary / 执行摘要", ("executive summary", "执行摘要", "管理摘要")),
        ("recommendations / 建议", ("recommendation", "recommendations", "建议", "推荐", "表达体系建议", "综合表达建议")),
        ("next steps / 下一步", ("next step", "next steps", "下一步", "行动计划", "后续工作")),
        ("limitations / 局限性", ("limitation", "limitations", "局限", "限制", "风险")),
    ]
    for label, markers in required:
        if not any(marker in text for marker in markers):
            warnings.append(f"Missing client-deliverable section: {label}")
    theoretical_markers = ("理论上", "一般来说", "in theory", "generally speaking", "通常情况下")
    if sum(text.count(marker) for marker in theoretical_markers) >= 3:
        warnings.append("Report may be too theoretical; replace generic background with client-specific findings and actions.")
    if markdown.count("|") < 6:
        warnings.append("Report has few or no tables; CRO consulting deliverables usually need concise evidence/recommendation tables.")
    if len(markdown) > 3000 and not re.search(r"(outputs/|文件|filename|sample|样本|数据)", markdown, re.IGNORECASE):
        warnings.append("Report does not clearly reference reviewed data, sample IDs, or generated output files.")
    dense_paragraphs = [
        para for para in re.split(r"\n\s*\n", markdown)
        if len(para) > 450 and not para.lstrip().startswith(("|", "```"))
    ]
    if dense_paragraphs:
        warnings.append("Report contains dense long paragraphs; split them into bullets, short paragraphs, or client implication tables.")
    if re.search(r"(?m)^[A-Z*]{100,}$", markdown):
        warnings.append("Report contains unwrapped long protein/sequence strings; wrap at 60-80 characters and move details to appendix.")
    if re.search(r"[$]{2,}|★|⭐|星级", markdown):
        warnings.append("Report uses star ratings or symbolic cost markers; replace with defined numeric scores or clear level definitions.")
    if re.search(r"(�|□|\?\?\?|乱码)", markdown):
        warnings.append("Report may contain garbled characters; inspect links, diagrams, and copied source text before final delivery.")
    if not re.search(r"(客户含义|客户影响|业务含义|client implication|implication|建议动作|行动)", markdown, re.IGNORECASE):
        warnings.append("Report lacks explicit client implications or action mapping; add 'client implication' and 'recommended action' columns.")
    return warnings


def run(args: dict[str, Any], *, session_id: str) -> dict[str, Any]:
    start = time.time()
    rel_path = str(args.get("markdown_path", "")).strip()
    if not rel_path:
        return {"stdout": "", "stderr": "Missing markdown_path", "exit_code": 2, "duration_ms": 0}
    try:
        md_path = _resolve_session_path(session_id, rel_path)
        if not md_path.exists() or not md_path.is_file():
            return {"stdout": "", "stderr": f"Markdown file not found: {rel_path}", "exit_code": 1, "duration_ms": 0}
        markdown = md_path.read_text(encoding="utf-8")
        basename = _safe_basename(str(args.get("output_basename") or md_path.stem))
        title = str(args.get("title") or basename)
        outputs = _session_root(session_id) / "outputs"
        outputs.mkdir(parents=True, exist_ok=True)
        docx_path = outputs / f"{basename}.docx"
        pdf_path = outputs / f"{basename}.pdf"
        blocks = _parse_blocks(markdown)
        _render_docx(blocks, docx_path, title)
        _render_pdf(blocks, pdf_path, title)
        warnings = _report_quality_warnings(markdown)
    except Exception as exc:
        return {
            "stdout": "",
            "stderr": f"render_report failed: {exc}",
            "exit_code": 1,
            "duration_ms": int((time.time() - start) * 1000),
        }
    return {
        "stdout": (
            f"Rendered report files:\n- outputs/{docx_path.name}\n- outputs/{pdf_path.name}"
            + (
                "\n\nReport quality warnings for CRO client deliverable:\n- "
                + "\n- ".join(warnings)
                + "\nRevise the Markdown before the next render pass if any warning is relevant."
                if warnings
                else "\n\nReport quality check: no major CRO deliverable warnings."
            )
        ),
        "stderr": "",
        "exit_code": 0,
        "duration_ms": int((time.time() - start) * 1000),
    }


def review_run(args: dict[str, Any], *, session_id: str) -> dict[str, Any]:
    start = time.time()
    rel_path = str(args.get("pdf_path", "")).strip()
    if not rel_path:
        return {"stdout": "", "stderr": "Missing pdf_path", "exit_code": 2, "duration_ms": 0}
    try:
        pdf_path = _resolve_session_path(session_id, rel_path)
        if not pdf_path.exists() or not pdf_path.is_file():
            return {"stdout": "", "stderr": f"PDF file not found: {rel_path}", "exit_code": 1, "duration_ms": 0}
        images = _pdf_page_images(pdf_path, int(args.get("max_pages") or 2))
        feedback = _call_multimodal_layout_review(images, str(args.get("focus") or "report layout readability, heading hierarchy, table fit, spacing, and professional polish"))
        review_path = (_session_root(session_id) / "outputs" / f"{pdf_path.stem}_layout-review.md")
        review_path.write_text(feedback.strip() + "\n", encoding="utf-8")
    except Exception as exc:
        return {
            "stdout": "",
            "stderr": f"review_report_layout failed: {exc}",
            "exit_code": 1,
            "duration_ms": int((time.time() - start) * 1000),
        }
    return {
        "stdout": f"Layout review saved:\n- outputs/{review_path.name}\n\n{feedback}",
        "stderr": "",
        "exit_code": 0,
        "duration_ms": int((time.time() - start) * 1000),
    }


def _pdf_page_images(pdf_path: Path, max_pages: int) -> list[str]:
    import fitz

    max_pages = max(1, min(max_pages, 4))
    images: list[str] = []
    with fitz.open(pdf_path) as doc:
        for i in range(min(max_pages, len(doc))):
            page = doc.load_page(i)
            pix = page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5), alpha=False)
            png = pix.tobytes("png")
            images.append("data:image/png;base64," + base64.b64encode(png).decode("ascii"))
    return images


def _call_multimodal_layout_review(images: list[str], focus: str) -> str:
    if not images:
        return "No PDF pages were available for layout review."
    base_url = os.getenv("HERMES_GATEWAY_BASE_URL", "http://llm-router:8642/v1").rstrip("/")
    api_key = os.getenv("HERMES_GATEWAY_API_KEY", "local-gateway")
    content: list[dict[str, Any]] = [
        {
            "type": "text",
            "text": (
                "You are reviewing rendered report pages for layout quality. "
                "The target is a client-facing CRO consulting deliverable, not an academic/theoretical essay. "
                "Give concise, actionable feedback in Chinese. Focus on: "
                f"{focus}; executive readability; evidence tables; recommendation clarity; whether the report looks ready for a consulting client. "
                "Return: 1) major issues, 2) concrete Markdown/layout/content revisions, 3) whether another render pass is needed."
            ),
        }
    ]
    content.extend({"type": "image_url", "image_url": {"url": image}} for image in images)
    payload = {
        "model": "multimodal",
        "messages": [{"role": "user", "content": content}],
        "max_completion_tokens": 800,
        "temperature": 0.2,
        "stream": False,
    }
    headers = {"Authorization": f"Bearer {api_key}", "Content-Type": "application/json"}
    with httpx.Client(timeout=float(os.getenv("HERMES_GATEWAY_TIMEOUT", "1200"))) as client:
        response = client.post(f"{base_url}/chat/completions", headers=headers, json=payload)
        response.raise_for_status()
        data = response.json()
    choice = (data.get("choices") or [{}])[0]
    message = choice.get("message") or {}
    return str(message.get("content") or message.get("reasoning_content") or "").strip()
